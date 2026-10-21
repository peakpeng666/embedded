from pathlib import Path
import json, re, subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT=Path(__file__).resolve().parents[1]
DOCX=ROOT/'作业2_书面解答.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.header_distance=Cm(.6);sec.footer_distance=Cm(.7)
styles=doc.styles
for name,size,east,latin in [('Normal',11,'宋体','Calibri'),('Title',24,'黑体','Calibri'),
                            ('Subtitle',10.5,'宋体','Calibri'),('Heading 1',16,'黑体','Calibri'),
                            ('Heading 2',12,'黑体','Calibri'),('Caption',9,'宋体','Calibri')]:
    st=styles[name];st.font.name=latin;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
    st.paragraph_format.space_before=Pt(0 if name in ('Normal','Title') else 8)
    st.paragraph_format.space_after=Pt(5 if name=='Normal' else 8)
    st.paragraph_format.line_spacing=1.15
styles['Normal'].paragraph_format.widow_control=True
styles['Caption'].paragraph_format.space_before=Pt(4)
styles['Caption'].paragraph_format.space_after=Pt(9)
styles['Caption'].font.color.rgb=RGBColor.from_string('53616B')
styles['Title'].paragraph_format.space_after=Pt(6)
styles['Subtitle'].font.italic=False
styles['Caption'].font.bold=False
for border in list(styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
styles['Heading 1'].paragraph_format.keep_with_next=True
styles['Heading 2'].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='嵌入式系统作业 2'
doc.core_properties.subject='Cortex-M4 与 STM32F407 基础及 GPIO 工程'
doc.core_properties.author='彭俊杰'
doc.core_properties.keywords='STM32F407ZGT6 STM32L496 Cortex-M4 GPIO'

def p(text='',style=None):
    para=doc.add_paragraph(text,style)
    return para
def h(text,level=1):return doc.add_heading(text,level)
def page():doc.add_page_break()
def caption(text):return p(text,'Caption')
def pic(name,width=16.8,alt=''):
    para=p();para.paragraph_format.space_after=Pt(0)
    para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(ROOT/'配图'/name),width=Cm(width))
    shape._inline.docPr.set('descr',alt or name)
    return shape
def table(headers,rows,widths,size=10):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D9D9D9');borders.append(n)
    props.append(borders)
    for rowno,vals in enumerate([headers]+rows):
        row=t.rows[0] if rowno==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if rowno==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for i,(cell,text,w) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for e,val in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                n=OxmlElement('w:'+e);n.set(qn('w:w'),val);n.set(qn('w:type'),'dxa');margins.append(n)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EEF1' if rowno==0 else 'FFFFFF');pr.append(shade)
            cp=cell.paragraphs[0];cp.paragraph_format.space_after=Pt(0);cp.paragraph_format.space_before=Pt(0)
            cp.paragraph_format.line_spacing=1.08
            r=cp.add_run(str(text));r.font.size=Pt(size);r.bold=(rowno==0)
            if rowno==0 or i==0:cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p().paragraph_format.space_after=Pt(2)
    return t
def code(text,size=9.3):
    # Code is native text; explicit spaces survive editing and PDF conversion.
    lines=text.strip('\n').splitlines()
    for i,line in enumerate(lines):
        para=p();fmt=para.paragraph_format
        fmt.space_after=Pt(0);fmt.space_before=Pt(0);fmt.line_spacing=Pt(11.7)
        fmt.left_indent=Cm(.2);fmt.keep_with_next=i<len(lines)-1
        for token in re.split(r'(\b(?:static|void|int|while|if|else|return|uint32_t|GPIO_InitTypeDef|GPIO_PinState)\b)',line):
            r=para.add_run(token);r.font.name='Consolas';r.font.size=Pt(size)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')
            if re.fullmatch(r'(static|void|int|while|if|else|return|uint32_t|GPIO_InitTypeDef|GPIO_PinState)',token):
                r.font.color.rgb=RGBColor.from_string('236A73')
    p().paragraph_format.space_after=Pt(0)
def link(para,label,url):
    rid=para.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'236A73');pr.append(c);run.append(pr)
    tx=OxmlElement('w:t');tx.text=label;run.append(tx);el.append(run);para._p.append(el)

# 1 — vocabulary. Each requested item appears explicitly.
p('嵌入式系统作业 2','Title')
p('彭俊杰　智装备241　2024306100324','Subtitle')
h('第1题 英文缩写')
terms=[
 ('ARM','Advanced RISC Machines','处理器体系结构与内核技术名称；早期源于 Acorn RISC Machine。'),
 ('AHB','Advanced High-performance Bus','AMBA 高性能总线，连接处理器、存储器和高速外设。'),
 ('AHB-Lite','AHB-Lite','AHB 的单主设备简化接口，省去多主设备仲裁信号。'),
 ('APB','Advanced Peripheral Bus','面向低带宽外设的简单总线，常经 AHB–APB 桥接入。'),
 ('PPB','Private Peripheral Bus','内核私有外设总线，访问内核控制与调试部件。'),
 ('DAP','Debug Access Port','调试访问端口，通过 JTAG 或 SWD 访问目标系统。'),
 ('JTAG','Joint Test Action Group','IEEE 1149.1 边界扫描接口，也用于芯片调试。'),
 ('SWD','Serial Wire Debug','串行线调试，核心信号为 SWDIO 和 SWCLK。'),
 ('MPU','Memory Protection Unit','存储器保护单元，设置区域访问权限和属性。'),
 ('FPU','Floating-Point Unit','浮点运算单元；Cortex-M4F 支持单精度浮点。'),
 ('NVIC','Nested Vectored Interrupt Controller','嵌套向量中断控制器，管理中断使能、挂起和优先级。'),
 ('ISR','Interrupt Service Routine','中断服务程序，响应并处理对应事件。'),
 ('Bitband','Bit-banding','位带机制，将一个位映射到别名区的一个 32 位字。'),
 ('SCS','System Control Space','系统控制空间，包含 SCB、NVIC、SysTick 等寄存器。'),
 ('SysTick','System Tick Timer','内核 24 位递减定时器，可产生周期性时基中断。'),
 ('MSR','Move to Special Register','写特殊寄存器的指令，如 MSR MSP, R0；MRS 为读取。'),
 ('MSP','Main Stack Pointer','主栈指针；复位后使用，异常处理模式固定使用。'),
 ('PC','Program Counter','程序计数器，即 R15，表示指令执行位置。'),
 ('SP','Stack Pointer','栈指针，即 R13；线程模式可选择 MSP 或 PSP。'),
 ('GPIO','General-Purpose Input/Output','通用输入输出端口，用于数字电平输入和输出。'),
 ('TIM','Timer','定时器，可用于计时、计数、输入捕获和 PWM。'),
 ('UART','Universal Asynchronous Receiver/Transmitter','通用异步收发器；USART 还支持同步模式。'),
 ('CANBUS','Controller Area Network Bus','控制器局域网总线，支持多节点通信与逐位仲裁。'),
]
table(['缩写','英文全称或规范名称','含义'],terms,[1.8,6.0,9.0],9.2)
caption('依据 Cortex-M4 编程手册与 STM32F4 参考手册整理。[1, 2]')

# 2 — memory.
page();h('第2题 存储器映射与外设地址')
p('Cortex-M4 使用 32 位地址，统一编址范围为 0x00000000～0xFFFFFFFF，共 4 GB。程序、数据和外设寄存器都通过地址访问；下图给出体系结构的主要分区。[1]')
pic('01_存储器映射.png',alt='Cortex-M4 地址空间，GPIO和CAN位于外设区，SysTick位于系统区')
caption('图 1 Cortex-M4 统一地址空间')
table(['对象','寄存器地址或基址','所属区域'],[
 ('GPIOA','0x40020000','外设区 → AHB1'),
 ('GPIOF','0x40021400','外设区 → AHB1'),
 ('CAN1 / CAN2','0x40006400 / 0x40006800','外设区 → APB1'),
 ('SysTick','0xE000E010～0xE000E01C','系统区 → PPB → SCS'),
],[3.0,7.0,6.8])
p('STM32F407ZGT6 的 1 MB Flash 位于 0x08000000～0x080FFFFF；SRAM1 与 SRAM2 合计 128 KB，位于 0x20000000～0x2001FFFF；64 KB CCM 数据 RAM 位于 0x10000000～0x1000FFFF，仅供 CPU 通过 D-bus 访问。[2]')
h('位带地址的含义',2)
p('SRAM 地址区的前 1 MB 和外设区的前 1 MB 分别具有位带别名区，基址为 0x22000000 和 0x42000000。别名地址 = 别名区基址 + 字节偏移 × 32 + 位号 × 4。对别名字进行读写，就能访问对应的单个位。[2]')

# 3 — exceptions.
page();h('第3题 中断概念与异常向量表')
p('中断是外设或外部事件发出的处理请求。处理器在满足响应条件后暂停当前程序，转入中断服务程序，结束后恢复执行。Cortex-M4 将外部中断纳入异常机制，与故障、SVC、PendSV 和 SysTick 等内核异常统一处理。[1]')
pic('02_中断流程.png',alt='中断请求经过挂起仲裁、硬件压栈、ISR处理和异常返回')
caption('图 2 中断响应过程')
p('关键词：中断源、使能、屏蔽、挂起、优先级、抢占、嵌套、向量、ISR、异常返回。优先级数值越小，优先级越高；高抢占优先级的异常可以打断低优先级的处理程序。外设标志通常由 ISR 按该外设规则清除。[1, 2]')
h('向量表的作用',2)
p('向量表把异常编号映射到处理函数入口。第 0 项保存初始 MSP，之后各项保存异常入口或保留值；每项占 4 字节。外部 IRQn 的向量偏移为 4 × (16 + IRQn)。VTOR 指定运行期间的向量表基址。[1]')
table(['表项','偏移','内容'],[
 ('0','0x00','初始主栈指针 __initial_sp'),('1','0x04','Reset_Handler'),
 ('2 / 3','0x08 / 0x0C','NMI_Handler / HardFault_Handler'),
 ('11 / 14','0x2C / 0x38','SVC_Handler / PendSV_Handler'),
 ('15','0x3C','SysTick_Handler'),('16 + IRQn','0x40 起','外部中断入口，例如 WWDG_IRQHandler'),
],[2.4,3.7,10.7])
p('基本异常栈帧包含 R0～R3、R12、LR、PC 和 xPSR。STM32F407 的外部中断编号为 IRQ0～IRQ81，因此本工程启动文件含 16 个系统表项和 82 个外部中断表项，共 98 项。[1, 2, 7]')

# 4 — buses.
page();h('第4题 STM32F407 片上总线')
p('STM32F407 通过 32 位多层 AHB 总线矩阵连接处理器、DMA、存储器和外设。主端口发起传输，从端口响应访问；RM0090 列出 8 个主端口和 7 个从端口。[2，第59～60页]')
pic('03_总线结构.png',alt='STM32F407总线矩阵的八个主端口和七个从端口')
caption('图 3 STM32F407 总线矩阵及端口')
h('题目指定外设的连接位置',2)
table(['模块','挂接总线','基址与用途'],[
 ('GPIO','AHB1','GPIOA 基址 0x40020000；通用输入输出'),
 ('TIM2','APB1','0x40000000；32 位通用定时器'),
 ('USART1','APB2','0x40011000；同步或异步串行通信'),
],[3.0,3.0,10.8])
p('AHB1 侧还包含 RCC、CRC、DMA 等模块，并通过两座桥连接 APB1 和 APB2；AHB2 连接 USB OTG FS、DCMI、RNG 等模块；外部存储器经 FSMC 接入。开启某个 GPIO 端口前，需要先打开它在 RCC 中的 AHB1 时钟使能位。[2]')

# 5 — clocks.
page();h('第5题 时钟源与系统时钟')
table(['类别','时钟源','频率与用途'],[
 ('高速内部','HSI','16 MHz RC；复位后的系统时钟来源'),
 ('高速外部','HSE','4～26 MHz 晶体；本例采用 8 MHz'),
 ('低速内部','LSI','标称约 32 kHz RC；可供独立看门狗、RTC 使用'),
 ('低速外部','LSE','32.768 kHz 晶体；常用作 RTC 时钟'),
 ('倍频输出','主 PLL','输入选择 HSI 或 HSE；经 M、N、P 得到 PLLCLK'),
],[3.0,2.7,11.1])
p('SYSCLK 可以直接选择 HSI、HSE，或选择主 PLL 输出。采用 PLL 时，SYSCLK = f输入 ÷ M × N ÷ P。例如 8 MHz ÷ 8 × 336 ÷ 2 = 168 MHz。[2, 3]')
pic('04_时钟树.png',alt='8MHz外部时钟经PLL得到168MHz，APB1为42MHz，APB2为84MHz')
caption('图 4 168 MHz 配置示例')
table(['时钟','最高频率','示例配置'],[
 ('SYSCLK','168 MHz','PLL 输出 168 MHz'),('AHB1 / HCLK','168 MHz','SYSCLK ÷ 1'),
 ('APB1 / PCLK1','42 MHz','HCLK ÷ 4'),('APB2 / PCLK2','84 MHz','HCLK ÷ 2'),
],[5.0,4.1,7.7])
p('配置顺序为开启时钟源、等待就绪、设置 PLL 和总线分频、设置供电档位与 Flash 等待周期，再切换 SYSCLK 并检查状态。以 2.7～3.6 V 供电运行在 168 MHz 时，Flash 配置 5 个等待周期。[2，表11]')
p('定时器时钟还需区分：APB 分频系数为 1 时，TIM 时钟等于 PCLK；大于 1 时，TIM 时钟为 2 × PCLK。因此本例 APB1 定时器为 84 MHz，APB2 定时器为 168 MHz。[2]')

# 6 — reset.
page();h('第6题 复位引脚与典型电路')
p('STM32F407ZGT6 采用 LQFP144 封装，复位引脚为第 25 脚 NRST。NRST 中的 N 表示低电平有效，RST 表示复位。封装顶视图中先根据定位标记找到 1 脚，再沿引脚编号定位 25 脚。[3，表7]')
pic('05_复位电路.png',alt='NRST第25脚通过10k电阻上拉到3.3V，并接100nF电容和复位按键到地')
caption('图 5 NRST 典型复位电路')
p('图中 10 kΩ 电阻将 NRST 上拉到 3.3 V，100 nF 电容接地，按键与电容并联。按下按键时 NRST 被拉低，芯片进入复位；松开后 NRST 恢复高电平，处理器按启动配置读取初始栈指针和复位入口。普中 F407-T100 原理图采用这一连接形式。[3, 5]')
table(['元件或信号','作用'],[
 ('10 kΩ 上拉电阻','让复位脚在按键松开时保持高电平。'),
 ('100 nF 电容','与上拉网络配合抑制复位脚上的短暂干扰。'),
 ('常开按键','闭合时将 NRST 接地，实现手动复位。'),
 ('调试器 NRST','连接同一复位网络，允许调试器控制复位。'),
],[4.2,12.6])
h('复位的其他来源',2)
p('除了 NRST 外部复位，芯片还支持上电或掉电复位、欠压复位、独立看门狗复位、窗口看门狗复位和软件复位。RCC_CSR 中的复位标志可用于判断复位来源。[2]')
p('芯片内部已有弱上拉电阻，数据手册给出的典型值为 40 kΩ。官方推荐复位脚外接 0.1 μF 电容；图中的外部 10 kΩ 上拉对应开发板的实际方案。[3，第121页]')

# 7 — chips.
page();h('第7题 主芯片数据手册与特征')
p('STM32F407ZGT6 侧重较高运算性能和丰富的高速接口；STM32L496 侧重低功耗与模拟、通信功能。两者均采用 Cortex-M4 内核，带单精度 FPU 和 DSP 指令。[3, 4]')
table(['项目','STM32F407ZGT6','STM32L496VET6'],[
 ('官方手册','DS8626 Rev 12\nSTM32F405xx / STM32F407xx','DS11585 Rev 20\nSTM32L496xx'),
 ('CPU 最高频率','168 MHz','80 MHz'),
 ('Flash','1 MB','512 KB'),
 ('SRAM','192 KB，其中 CCM 为 64 KB；另有 4 KB 备份 SRAM','320 KB，其中 64 KB 带硬件奇偶校验'),
 ('封装','LQFP144','LQFP100'),
 ('工作电源','1.8～3.6 V','1.71～3.6 V'),
 ('主要接口','以太网 MAC、USB OTG FS/HS、CAN、USART、SPI、I²C、FSMC','USB OTG FS、CAN、LPUART、I²C、SPI、QUADSPI、SDMMC'),
 ('模拟与控制','3 个 12 位 ADC、2 个 DAC 通道、多组 TIM','3 个 12 位 ADC、2 个 DAC 通道、运放、比较器、低功耗定时器'),
 ('典型特点','较高主频，适合控制、图像采集和联网应用','多种低功耗模式，适合电池供电和物联网终端'),
],[2.9,6.95,6.95],10)
p('L496 系列最高支持 1 MB Flash，表中按 VET6 的实际容量列为 512 KB。F407 的以太网模块是 MAC；外部网络连接还需要 PHY。USB HS 的高速通信需要外部 ULPI PHY。[3, 4]')
h('开发板与芯片型号的对应',2)
p('本地普中 F407-T100 原理图标注主芯片为 STM32F407ZGT6。RT-Thread 与正点原子公开的常见潘多拉板使用 STM32L475VET6，参数为 80 MHz、512 KB Flash、128 KB RAM；它与题目列出的 STM32L496 是不同型号。[5, 6]')
p('两份官方 PDF 已随作业放入“参考资料”目录：STM32F407ZGT6数据手册_ST.pdf、STM32L496数据手册_ST.pdf。')

# 8.1 — schematic.
page();h('第8题 扩展练习一 最小系统电路')
p('STM32F407ZGT6 的基本运行电路包括电源及去耦、复位、启动配置和调试接口。下图采用 3.3 V 供电与 8 MHz 外部晶振，并结合官方硬件指南与普中原理图重绘。[3, 5, 8]')
pic('06_最小系统.png',alt='STM32F407ZGT6最小系统，包含供电、VCAP电容、8MHz晶振、BOOT0、NRST与SWD')
caption('图 6 STM32F407ZGT6 最小系统连接示意')
p('VDD 与 VDDA 接 3.3 V，VSS 与 VSSA 接地。每个 VDD 附近放置 100 nF 去耦电容；VCAP1、VCAP2 分别接 2.2 μF 电容到地，为片内调压器稳定工作提供条件。PDR_ON 接高电平，使内部复位电路工作。[3, 8]')
p('BOOT0 下拉时从用户 Flash 启动。PA13 和 PA14 分别用于 SWDIO 与 SWCLK；调试器同时连接 GND、3.3 V 参考和 NRST。RTC 若需要外部低速时钟，可在 PC14、PC15 之间增加 32.768 kHz 晶体及匹配电容。[2, 3]')
p('完整开发板原理图为“普中F407-T100开发板原理图_STM32F407ZGT6.pdf”，随作业附在“参考资料”目录，可对照其中 CPU、POWER、RESET、BOOT、JTAG 模块查看具体元件编号。[5]')

# 8.2 and 8.3 — project and startup.
page();h('第8题 扩展练习二 工程模板')
p('标准库与 HAL 工程都由启动文件、CMSIS、库驱动和用户代码组成。Group 是 µVision 中的源文件分组；main.c 组织初始化与主循环。模板预先确定芯片、头文件路径、宏定义、存储器布局和编译器设置。')
p('两类目录中的 Blank.uvprojx 是空白模板，入口文件为 User/main_blank.c；Template.uvprojx 是 GPIO 示例，入口文件为 User/main.c。')
table(['组成','标准库工程','HAL 工程'],[
 ('用户程序','User/main.c 与中断处理文件','User/main.c 与中断处理文件'),
 ('启动与内核','startup_stm32f40_41xxx.s\nCMSIS 与 system_stm32f4xx.c','startup_stm32f407xx.s\nCMSIS 与 system_stm32f4xx.c'),
 ('驱动库','STM32F4xx_StdPeriph_Driver','STM32F4xx_HAL_Driver'),
 ('空白模板入口','main → while (1)','HAL_Init → SystemClock_Config\n→ while (1)'),
],[2.9,6.95,6.95],9.8)
h('第8题 扩展练习三 启动文件')
p('实际 SPL 工程的启动文件设置 Stack_Size = 0x00000400，即 1 KB 栈；Heap_Size = 0x00000200，即 512 B 堆。SPACE 保留空间，DCD 放置 32 位数据或地址。以下节选自工程中的汇编文件。[7]')
code('''__Vectors       DCD     __initial_sp
                DCD     Reset_Handler
                DCD     NMI_Handler
                DCD     HardFault_Handler

Reset_Handler   PROC
                EXPORT  Reset_Handler [WEAK]
                IMPORT  SystemInit
                IMPORT  __main
                LDR     R0, =SystemInit
                BLX     R0
                LDR     R0, =__main
                BX      R0
                ENDP''',9.5)
pic('07_启动流程.png',alt='复位后装载MSP，经过Reset_Handler、SystemInit和C运行时进入main')
caption('图 7 从复位到 main 的执行顺序')
p('硬件从启动映射区读取初始 MSP 和复位向量。SystemInit 完成系统级初始化；__main 建立 C 运行环境，然后调用 main。启动文件把大部分异常处理函数定义为弱符号，用户定义同名函数后即可替换默认入口。[1, 7]')

# 8.4 SPL.
page();h('第8题 扩展练习四 标准库 GPIO 工程')
p('两份工程使用相同的引脚连接：PF9 驱动低电平点亮的 LED，PE4 连接按下接地的按键。程序先开启端口时钟，配置 PF9 为推挽输出、PE4 为上拉输入，再循环读取按键并控制 LED。[5, 7]')
caption('代码 1 SPL_GPIO/User/main.c 完整内容')
code((ROOT/'_build/SPL_display.c').read_text(encoding='utf-8'),9.1)
p('GPIO_SetBits 将输出置高，GPIO_ResetBits 将输出置低；GPIO_ReadInputDataBit 读取输入数据位。GPIO_StructInit 先填入默认配置，再修改本实验需要的字段。')
p('程序逻辑为：按下按键，PE4 读到 0，将 PF9 清零，LED 点亮；松开按键，PE4 由上拉保持为 1，将 PF9 置位，LED 熄灭。')

# 8.4 HAL.
page();h('第8题 扩展练习四 HAL GPIO 工程')
p('HAL 工程先调用 HAL_Init，再配置系统时钟与 GPIO。下面是实际 main.c 中的 GPIO 初始化函数和主循环；SystemClock_Config 的完整实现随工程提供，使用 HSE 8 MHz 与 PLL 得到 168 MHz。')
caption('代码 2 HAL_GPIO/User/main.c 中的初始化与主循环')
code((ROOT/'_build/HAL_display.c').read_text(encoding='utf-8'),9.2)
p('HAL_GPIO_ReadPin 返回 GPIO_PIN_RESET 或 GPIO_PIN_SET；将这个状态写入 PF9，即可使低电平 LED 跟随低有效按键。初始化时先写高电平，再将 PF9 切换为输出，使 LED 初始熄灭。')
table(['操作','标准库','HAL'],[
 ('打开时钟','RCC_AHB1PeriphClockCmd','__HAL_RCC_GPIOx_CLK_ENABLE'),
 ('初始化','GPIO_Init','HAL_GPIO_Init'),
 ('读输入','GPIO_ReadInputDataBit','HAL_GPIO_ReadPin'),
 ('写输出','GPIO_SetBits / GPIO_ResetBits','HAL_GPIO_WritePin'),
],[2.7,7.05,7.05],9.2)

# Verification and references.
page();h('工程编译记录与参考资料')
p('两份空白模板和两份 GPIO 工程已在本机使用 Keil µVision 和 Arm Compiler 5.06 update 7 build 960 重新编译。工程内包含源码、µVision 工程文件、HEX、AXF 和编译日志。')
results=json.loads((ROOT/'GPIO工程/build_results.json').read_text(encoding='utf-8'))
table(['工程','编译结果','代码与只读数据'],[
 ('SPL_Blank','0 错误，0 警告','Code 704 B，RO-data 408 B'),
 ('HAL_Blank','0 错误，0 警告','Code 3106 B，RO-data 450 B'),
 ('SPL_GPIO','0 错误，0 警告','Code 1068 B，RO-data 408 B'),
 ('HAL_GPIO','0 错误，0 警告','Code 3814 B，RO-data 450 B'),
],[3.4,5.2,8.2])
p('本次完成源码与编译验证；按键控制 LED 的现象由接线和代码推导，尚未进行实板测试。')
h('参考资料',2)
refs=[
 ('[1] STMicroelectronics. PM0214 Rev 10. Cortex-M4 编程手册。用于存储器映射、特殊寄存器、异常机制与 SysTick。','ST 官方 PDF','https://www.st.com/resource/en/programming_manual/pm0214-stm32-cortexm4-mcus-and-mpus-programming-manual-stmicroelectronics.pdf'),
 ('[2] STMicroelectronics. RM0090 Rev 22. STM32F405/407 等器件参考手册。第2章总线与存储器；第6章 RCC；GPIO、TIM、USART、CAN 章节。','ST 参考手册入口','https://www.st.com/en/microcontrollers-microprocessors/stm32f407zg.html#documentation'),
 ('[3] STMicroelectronics. DS8626 Rev 12，2026年3月。STM32F405xx / STM32F407xx 数据手册。表7为引脚表，第121页为 NRST 电气特性。','ST 官方 PDF','https://www.st.com/resource/en/datasheet/stm32f407zg.pdf'),
 ('[4] STMicroelectronics. DS11585 Rev 20，2025年10月。STM32L496xx 数据手册；VET6 容量另参见产品页。','ST 官方 PDF','https://www.st.com/resource/en/datasheet/stm32l496qe.pdf'),
 ('[5] 普中 F407-T100 开发板原理图，图纸日期 2020年12月24日。主芯片为 STM32F407ZGT6。本地课程资料副本，完整 PDF 随作业附后。',None,None),
 ('[6] RT-Thread 官方仓库。STM32L475 潘多拉开发板 BSP 说明。用于核对常见潘多拉板的 MCU 与存储器容量。','官方 BSP 说明','https://github.com/RT-Thread/rt-thread/blob/master/bsp/stm32/stm32l475-atk-pandora/README.md'),
 ('[7] 本次 GPIO工程/SPL_GPIO 与 GPIO工程/HAL_GPIO。标准库模板取自本地普中例程，HAL 模板取自本地 HAL 最小工程；main.c 为本次 GPIO 练习代码。',None,None),
 ('[8] STMicroelectronics. AN4488. Getting started with STM32F4xxxx MCU hardware development。用于电源、复位、启动及调试电路。','ST 官方 PDF','https://www.st.com/resource/en/application_note/an4488-getting-started-with-stm32f4xxxx-mcu-hardware-development-stmicroelectronics.pdf'),
]
for text,label,url in refs:
    para=p(text);para.paragraph_format.space_after=Pt(7)
    for r in para.runs:r.font.size=Pt(10)
    if url:
        para.add_run(' ');link(para,label,url)

doc.save(DOCX)
print(DOCX)
